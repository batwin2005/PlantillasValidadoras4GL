# -*- coding: utf-8 -*-
"""
Procesador de plantillas Excel genérico usando openpyxl.
Usa celdas con etiquetas {{tag}} o rangos nombrados como marcadores de posición.
Permite reemplazar valores por etiquetas manteniendo formato, estilos, validaciones, etc.
"""

import openpyxl
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils.cell import coordinate_from_string
from openpyxl.utils import column_index_from_string
from copy import copy as copy_func
from pathlib import Path
from typing import Dict, Any, List, Optional, Union
from copy import copy as copy_func
import re

class ExcelTemplateProcessor:
    """
    Procesador de plantillas Excel genérico.
    
    Uso:
    1. Crear plantilla Excel con celdas que tengan etiquetas {{tag}} o rangos nombrados
    2. Los nombres de las celdas/rangos son las "etiquetas" (tags)
    4. Genera un nuevo archivo con los valores reemplazados
    
    Ventajas:
    - Usa Excel nativo (.xlsx) - se abre en cualquier Excel
    - Mantiene formato, estilos, validaciones, fórmulas, etc.
    - Funciona con cualquier plantilla existente
    - Etiquetas legibles (ej: "cliente.nombre", "cuenta.saldo")
    - Soporta rangos (para tablas/listas) y celdas individuales
    """
    
    def __init__(self, template_path: Union[str, Path]):
        self.template_path = Path(template_path)
        self.workbook = None
        self.tag_map = {}  # tag -> (worksheet, cell_ref o range)
        
    def load_template(self) -> 'ExcelTemplateProcessor':
        """Carga la plantilla y descubre las etiquetas (nombres de celdas/rangos)."""
        self.workbook = openpyxl.load_workbook(self.template_path)
        self._discover_tags()
        return self
    
    def _discover_tags(self):
        """Descubre etiquetas de celdas con {{tag}} y rangos nombrados."""
        self.tag_map = {}
        import re
        
        # 1. Rangos nombrados (Named Ranges) - parsear attr_text directamente
        for name, named_range in self.workbook.defined_names.items():
            try:
                attr = named_range.attr_text  # ej: "Hoja!A1:B10" o "'Hoja Con Espacios'!A1:B10"
                if attr:
                    # Parsear: [hoja!]rango
                    if '!' in attr:
                        sheet_part, range_part = attr.split('!', 1)
                        # Quitar comillas si las hay
                        sheet_name = sheet_part.strip("'")
                        ws = self.workbook[sheet_name]
                        self.tag_map[name] = ('range', ws, range_part)
            except Exception as e:
                print(f"Warning: Named Range '{name}' no procesado: {e}")
                pass
        
        # 2. Celdas con etiquetas {{tag}} - tanto tags completos como embebidos
        for ws in self.workbook.worksheets:
            for row in ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=False):
                for cell in row:
                    if cell.value and isinstance(cell.value, str):
                        cell_str = cell.value
                        
                        # Tag completo: "{{tag}}"
                        if cell_str.startswith('{{') and cell_str.endswith('}}'):
                            tag = cell_str[2:-2].strip()
                            self.tag_map[tag] = ('cell', ws, cell.coordinate)
                        
                        # Tags embebidos: "texto {{tag}} mas texto"
                        elif '{{' in cell_str and '}}' in cell_str:
                            matches = re.findall(r'\{\{(\w+)\}\}', cell_str)
                            for tag in matches:
                                self.tag_map[tag] = ('cell_embedded', ws, cell.coordinate, cell_str)
        
        # 3. Celdas con comentario que indique etiqueta {{tag}}
        for ws in self.workbook.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    if cell.comment and cell.comment.text:
                        matches = re.findall(r'\{\{(\w+)\}\}', cell.comment.text)
                        for tag in matches:
                            if tag not in self.tag_map:
                                self.tag_map[tag] = ('cell', ws, cell.coordinate)
    
    def get_available_tags(self) -> Dict[str, str]:
        """Retorna diccionario de etiquetas disponibles y su tipo."""
        result = {}
        for tag, tag_info in self.tag_map.items():
            tipo = tag_info[0]
            result[tag] = tipo
        return result
    
    def set_value(self, tag: str, value: Any) -> bool:
        """Establece un valor en una etiqueta."""
        if tag not in self.tag_map:
            return False
        
        tag_info = self.tag_map[tag]
        tipo = tag_info[0]
        ws = tag_info[1]
        ref = tag_info[2]
        
        if tipo == 'cell':
            cell = ws[ref]
            cell.value = value
        elif tipo == 'cell_embedded':
            # Tag embebido en texto: reemplazar {{tag}} en el contenido
            cell = ws[ref]
            original_text = tag_info[3]
            tag_pattern = f'{{{{{tag}}}}}'
            cell.value = original_text.replace(tag_pattern, str(value))
        elif tipo == 'range':
            # Para rangos, escribir en la celda superior izquierda
            cell = ws[ref.split(':')[0]] if ':' in ref else ws[ref]
            cell.value = value
        return True
    
    def set_values(self, values: Dict[str, Any]) -> Dict[str, bool]:
        """Establece múltiples valores. Retorna dict de resultados."""
        results = {}
        for tag, value in values.items():
            results[tag] = self.set_value(tag, value)
        return results
    
    def set_row_data(self, tag: str, rows_data: List[Dict], start_row: int = None) -> bool:
        """
        Llena un rango con múltiples filas de datos (para tablas).
        tag: etiqueta del rango
        rows_data: lista de diccionarios {columna: valor}
        start_row: fila inicial (si None, usa la fila del rango)
        """
        if tag not in self.tag_map:
            return False
        
        tipo, ws, ref = self.tag_map[tag]
        if tipo != 'range':
            return False
        
        # Parsear rango: ej "Hoja!A2:D10" -> start_cell="A2", end_cell="D10"
        if ':' in ref:
            start_cell, end_cell = ref.split(':')
        else:
            start_cell = ref
            end_cell = ref
        
        # Quitar nombre de hoja si existe (ej: "Hoja!A9" -> "A9")
        for c in [start_cell, end_cell]:
            if '!' in c:
                c = c.split('!', 1)[1]
        start_cell = start_cell.split('!', 1)[1] if '!' in start_cell else start_cell
        end_cell = end_cell.split('!', 1)[1] if '!' in end_cell else end_cell
        
        # Obtener fila y columna inicial
        start_col_letter, start_row_num = coordinate_from_string(start_cell)
        end_col_letter, end_row_num = coordinate_from_string(end_cell)
        
        start_col = column_index_from_string(start_col_letter)
        end_col = column_index_from_string(end_col_letter)
        
        if start_row is None:
            start_row = start_row_num
        
        # Obtener encabezados de la primera fila del rango
        headers = []
        header_keys = []  # Claves reales para buscar en row_data
        for col_idx in range(start_col, end_col + 1):
            header_cell = ws.cell(row=start_row_num, column=col_idx)
            if header_cell.value:
                header_value = str(header_cell.value).strip()
                headers.append(header_value)
                # Extraer clave real del tag {{campo}}
                if header_value.startswith('{{') and header_value.endswith('}}'):
                    header_keys.append(header_value[2:-2].strip())
                else:
                    header_keys.append(header_value)
            else:
                headers.append(f'col_{col_idx}')
                header_keys.append(f'col_{col_idx}')
        
        # Escribir filas de datos (empezando DESPUÉS de los headers)
        data_start_row = start_row + 1  # Saltar fila de headers
        for i, row_data in enumerate(rows_data):
            row_num = data_start_row + i
            for col_idx, header in enumerate(headers):
                col = start_col + col_idx
                if col <= end_col:
                    cell = ws.cell(row=row_num, column=col)
                    # Usar la clave real para buscar en row_data
                    key = header_keys[col_idx]
                    cell.value = row_data.get(key, '')
                    # Copiar estilo de la fila de encabezado si existe
                    header_cell = ws.cell(row=start_row_num, column=col_idx + start_col)
                    if header_cell.has_style:
                        cell._style = copy_func(header_cell._style)
        
        return True
    
    def save(self, output_path: Union[str, Path]) -> 'ExcelTemplateProcessor':
        """Guarda el workbook procesado."""
        output_path = Path(output_path)
        self.workbook.save(output_path)
        return self
    
    def close(self):
        """Cierra el workbook."""
        if self.workbook:
            self.workbook.close()

    def auto_tag_labels(self, label_col: int = 1, value_col: int = 2, 
                        prefix: str = "cliente", skip_keywords: list = None) -> int:
        """
        Auto-genera tags {{prefix.campo}} en value_col basandose en etiquetas de label_col.
        Detecta filas donde label_col tiene texto y value_col esta vacio o es placeholder.
        
        Args:
            label_col: columna con etiquetas (1=A, 2=B, ...)
            value_col: columna donde poner tags (1=A, 2=B, ...)
            prefix: prefijo del tag (ej: "cliente" -> {{cliente.campo}})
            skip_keywords: palabras en etiqueta para saltar (secciones, titulos)
            
        Returns:
            numero de tags creados
        """
        if self.workbook is None:
            self.load_template()
        
        if skip_keywords is None:
            skip_keywords = ['instrucciones', 'formulario', 'plantilla', 'importante',
                           'datos personales', 'datos de contacto', 'informacion', 
                           'validacion', 'observaciones', 'seccion']
        
        import re
        def slugify(text):
            text = text.lower().strip()
            text = text.replace('ñ', 'n').replace('á', 'a').replace('é', 'e').replace('í', 'i').replace('ó', 'o').replace('ú', 'u')
            text = re.sub(r'[^a-z0-9]+', '_', text)
            text = re.sub(r'_+', '_', text).strip('_')
            return text
        
        placeholder_words = ['escriba', 'seleccione', 'use', 'aaaa', 'cliente@', 
                           'fictici', 'ejemplo', 'opcion', 'datos', '---']
        
        created = 0
        for ws in self.workbook.worksheets:
            for row in ws.iter_rows(min_row=1, max_row=ws.max_row,
                                    min_col=label_col, max_col=value_col, values_only=False):
                label_cell = row[0]
                value_cell = row[1] if len(row) > 1 else None
                
                if not label_cell.value or not isinstance(label_cell.value, str):
                    continue
                label = label_cell.value.strip()
                if len(label) < 3 or len(label) > 60:
                    continue
                
                # Saltar secciones/titulos
                if re.match(r'^\d+\.', label): continue
                if any(kw in label.lower() for kw in skip_keywords): continue
                
                # Verificar si value_col esta vacio o es placeholder
                if value_cell is None:
                    continue
                v = value_cell.value
                if v is None or v == '':
                    is_placeholder = True
                else:
                    v_str = str(v).lower()
                    is_placeholder = any(w in v_str for w in placeholder_words)
                
                if is_placeholder:
                    tag = f'{{{{{prefix}.{slugify(label)}}}}}'
                    value_cell.value = tag
                    created += 1
        
        # Re-descubrir tags
        self._discover_tags()
        return created


# ============================================================
# FUNCIONES DE CONVENIENCIA
# ============================================================

def procesar_plantilla(template_path: Union[str, Path], 
                       datos: Dict[str, Any], 
                       output_path: Union[str, Path],
                       filas_tabla: Optional[Dict[str, List[Dict]]] = None) -> bool:
    """
    Función de conveniencia para procesar una plantilla completa.
    
    Args:
        template_path: ruta a la plantilla .xlsx
        datos: dict {etiqueta: valor} para celdas individuales
        output_path: ruta del archivo de salida
        filas_tabla: dict {etiqueta_rango: lista_de_dicts} para tablas
    
    Returns:
        True si éxito, False si error
    """
    try:
        processor = ExcelTemplateProcessor(template_path)
        processor.load_template()
        
        # Valores simples
        if datos:
            processor.set_values(datos)
        
        # Tablas/datos tabulares
        if filas_tabla:
            for tag, filas in filas_tabla.items():
                processor.set_row_data(tag, filas)
        
        processor.save(output_path)
        processor.close()
        return True
    except Exception as e:
        print(f"Error procesando plantilla: {e}")
        return False


def crear_plantilla_ejemplo(output_path: Union[str, Path]) -> bool:
    """Crea una plantilla de ejemplo con etiquetas."""
    try:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Cliente"
        
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.worksheet.datavalidation import DataValidation
        from openpyxl.workbook.defined_name import DefinedName
        
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="003366", end_color="003366", fill_type="solid")
        header_align = Alignment(horizontal="center")
        thin_border = Border(
            left=Side(style='thin'), right=Side(style='thin'),
            top=Side(style='thin'), bottom=Side(style='thin')
        )
        
        ws = wb.active
        ws.title = "Cliente"
        
        etiquetas = [
            ("A1", "{{cliente.tipo_doc}}"), ("B1", "DNI"),
            ("A2", "{{cliente.num_doc}}"), ("B2", "30123456"),
            ("A3", "{{cliente.apellidos}}"), ("B3", "PEREZ"),
            ("A4", "{{cliente.nombres}}"), ("B4", "JUAN CARLOS"),
            ("A5", "{{cliente.fecha_nac}}"), ("B5", "03/15/1985"),
            ("A6", "{{cliente.telefono}}"), ("B6", "011-4567-8901"),
            ("A7", "{{cliente.email}}"), ("B7", "juan@email.com"),
            ("A8", "{{cliente.direccion}}"), ("B8", "AV. CORRIENTES 1234"),
            ("A9", "{{cliente.ciudad}}"), ("B9", "CABA"),
        ]
        
        for cell_ref, tag in etiquetas:
            cell = ws[cell_ref]
            cell.value = tag
            cell.font = Font(bold=True, color="003366")
        
        ws2 = wb.create_sheet("Cuentas")
        
        headers = ["{{cuenta.numero}}", "{{cuenta.tipo}}", "{{cuenta.moneda}}", 
                   "{{cuenta.saldo}}", "{{cuenta.fecha_apertura}}", "{{cuenta.estado}}"]
        for col_idx, tag in enumerate(headers, 1):
            cell = ws2.cell(row=1, column=col_idx, value=tag)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="003366", end_color="003366", fill_type="solid")
            cell.alignment = Alignment(horizontal="center")
            cell.border = Border(left=Side(style='thin'), right=Side(style='thin'),
                                 top=Side(style='thin'), bottom=Side(style='thin'))
        
        from openpyxl.workbook.defined_name import DefinedName
        table_range = DefinedName(name="tabla_cuentas", attr_text="Cuentas!A1:F20")
        wb.defined_names.add(table_range)
        
        dv_tipo = DataValidation(type="list", formula1='"AHORRO,CORRIENTE,PLAZOFIJO"', allow_blank=False)
        dv_tipo.error = "Seleccione tipo valido"
        dv_tipo.errorTitle = "Tipo Invalido"
        dv_tipo.sqref = "B2:B100"
        ws2 = wb["Cuentas"]
        ws2.add_data_validation(dv_tipo)
        
        wb.save(output_path)
        print("Plantilla de ejemplo creada en:", output_path)
        return True
    except Exception as e:
        print(f"Error creando plantilla: {e}")
        return False


if __name__ == "__main__":
    import sys
    
    # 1. Crear plantilla de ejemplo
    template_path = Path("plantilla_banco_ejemplo.xlsx")
    if not template_path.exists():
        crear_plantilla_ejemplo(template_path)
    
    # 2. Procesar plantilla con datos
    datos_cliente = {
        "cliente.tipo_doc": "DNI",
        "cliente.num_doc": "30123456",
        "cliente.apellidos": "PEREZ",
        "cliente.nombres": "JUAN CARLOS",
        "cliente.fecha_nac": "03/15/1985",
        "cliente.telefono": "011-4567-8901",
        "cliente.email": "juan.perez@email.com",
        "cliente.direccion": "AV. CORRIENTES 1234",
        "cliente.ciudad": "CABA",
    }
    
    # Datos para tabla de cuentas
    cuentas_data = [
        {"cuenta.numero": "0010000123456789", "cuenta.tipo": "AHORRO", 
         "cuenta.moneda": "ARS", "cuenta.saldo": 150000.50, 
         "cuenta.fecha_apertura": "2020-01-15", "cuenta.estado": "ACTIVA"},
        {"cuenta.numero": "0010000123456790", "cuenta.tipo": "CORRIENTE",
         "cuenta.moneda": "USD", "cuenta.saldo": 5000.00,
         "cuenta.fecha_apertura": "2021-03-10", "cuenta.estado": "ACTIVA"},
    ]
    
    filas_tabla = {"tabla_cuentas": cuentas_data}
    
    output_path = Path("cliente_procesado.xlsx")
    
    exito = procesar_plantilla(
        template_path=template_path,
        datos={"cliente.tipo_doc": "DNI", "cliente.num_doc": "30123456",
               "cliente.apellidos": "PEREZ", "cliente.nombres": "JUAN CARLOS",
               "cliente.fecha_nac": "03/15/1985", "cliente.telefono": "011-4567-8901",
               "cliente.email": "juan.perez@email.com", "cliente.direccion": "AV. CORRIENTES 1234",
               "cliente.ciudad": "CABA"},
        output_path=output_path,
        filas_tabla={"tabla_cuentas": cuentas_data}
    )
    
    if exito:
        print(f"\nProcesado exitoso! Archivo generado: {output_path}")
    else:
        print("Error en el procesamiento")
        sys.exit(1)