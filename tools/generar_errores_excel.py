#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
generar_errores_excel.py
Genera archivo Excel con hoja ERRORES a partir de validacion_detalle
Uso: python generar_errores_excel.py <resultado_id> <archivo_original> <archivo_salida>
"""

import sys
import os
import shutil
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

def get_errores(resultado_id):
    """
    En producción: conectar a Informix
    SELECT num_fila, columna_excel, nombre_campo, valor_original, 
           valor_normalizado, estado, codigo_error, mensaje, accion_sugerida
    FROM validacion_detalle 
    WHERE resultado_id = ? AND estado IN ('ERROR','WARNING','DUPLICADO')
    ORDER BY num_fila, campo_id
    """
    # MOCK
    return [
        (3, 'C', 'cliente.apellidos', 'PEREZ123', 'PEREZ123', 'ERROR', 'ERR_NOMBRE', 
         'Apellidos solo letras: PEREZ123', 'CORREGIR'),
        (3, 'F', 'cliente.email', 'juanemail.com', 'juanemail.com', 'ERROR', 'ERR_EMAIL',
         'Email inválido: juanemail.com', 'CORREGIR'),
        (5, 'B', 'cliente.num_doc', '30123456', '30123456', 'DUPLICADO', 'ERR_UNIQUE',
         'Cliente ya existe con este documento', 'OMITIR'),
    ]

def get_info_validacion(resultado_id):
    return {
        'plantilla': 'Clientes Banco Nación',
        'banco': 'Banco de la Nación Argentina',
        'archivo_original': 'datos_clientes.xlsx',
        'fecha': '12/10/2026 15:30',
    }

def generar_archivo_con_errores(resultado_id, archivo_original, archivo_salida):
    try:
        print(f"[INFO] Generando archivo con errores: {archivo_salida}")
        
        # 1. Copiar archivo original
        shutil.copy2(archivo_original, archivo_salida)
        
        # 2. Abrir y agregar hoja ERRORES
        wb = openpyxl.load_workbook(archivo_salida)
        
        # Eliminar hoja ERRORES si existe
        if 'ERRORES' in wb.sheetnames:
            del wb['ERRORES']
        
        ws_err = wb.create_sheet("ERRORES", 0)  # Primera hoja
        
        # Estilos
        header_font = Font(bold=True, color="FFFFFF", size=11)
        header_fill = PatternFill(start_color="CC0000", end_color="CC0000", fill_type="solid")
        header_align = Alignment(horizontal="center", wrap_text=True)
        title_font = Font(bold=True, size=14, color="CC0000")
        subtitle_font = Font(size=10, color="666666")
        error_fill = PatternFill(start_color="FFE6E6", end_color="FFE6E6", fill_type="solid")
        warning_fill = PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid")
        dup_fill = PatternFill(start_color="FFE0B2", end_color="FFE0B2", fill_type="solid")
        thin_border = Border(
            left=Side(style='thin'), right=Side(style='thin'),
            top=Side(style='thin'), bottom=Side(style='thin')
        )
        
        # Cabecera info
        info = get_info_validacion(resultado_id)
        ws_err.merge_cells('A1:I1')
        ws_err['A1'] = f"⚠️ ARCHIVO CON ERRORES - {info['banco']}"
        ws_err['A1'].font = title_font
        ws_err['A1'].alignment = Alignment(horizontal="center")
        
        ws_err.merge_cells('A2:I2')
        ws_err['A2'] = f"Plantilla: {info['plantilla']} | Archivo: {info['archivo_original']} | Fecha: {info['fecha']}"
        ws_err['A2'].font = subtitle_font
        ws_err['A2'].alignment = Alignment(horizontal="center")
        
        ws_err.merge_cells('A3:I3')
        ws_err['A3'] = "INSTRUCCIONES: Corrija los errores marcados en la hoja 'Datos' y vuelva a validar (Opción 4). No elimine esta hoja."
        ws_err['A3'].font = Font(italic=True, color="CC0000", size=10)
        ws_err['A3'].alignment = Alignment(horizontal="center", wrap_text=True)
        
        # Encabezados errores (fila 5)
        headers = ['Fila', 'Columna', 'Campo', 'Valor Original', 'Valor Normalizado', 
                   'Estado', 'Código Error', 'Mensaje', 'Acción Sugerida']
        row_h = 5
        for idx, h in enumerate(headers, 1):
            cell = ws_err.cell(row=row_h, column=idx, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align
            cell.border = thin_border
        
        # Datos errores
        errores = get_errores(resultado_id)
        fill_map = {'ERROR': error_fill, 'WARNING': warning_fill, 'DUPLICADO': dup_fill}
        
        for row_idx, (fila, col, campo, val_orig, val_norm, estado, cod, msg, accion) in enumerate(errores, row_h + 1):
            vals = [fila, col, campo, val_orig, val_norm, estado, cod, msg, accion]
            for col_idx, val in enumerate(vals, 1):
                cell = ws_err.cell(row=row_idx, column=col_idx, value=val)
                cell.border = thin_border
                cell.alignment = Alignment(wrap_text=True)
                
                # Colorear fila según estado
                if col_idx == 1:  # Columna Fila
                    fill = fill_map.get(estado, error_fill)
                    for c in range(1, 10):
                        ws_err.cell(row=row_idx, column=c).fill = fill
        
        # Ajustar anchos
        anchos = [8, 10, 25, 25, 25, 12, 18, 40, 18]
        for idx, w in enumerate(anchos, 1):
            ws_err.column_dimensions[get_column_letter(idx)].width = w
        
        # Auto-filtro
        ws_err.auto_filter.ref = f"A{row_h}:I{row_h + len(errores)}"
        ws_err.freeze_panes = f"A{row_h + 1}"
        
        # 3. También marcar errores en la hoja de datos original
        if 'Datos' in wb.sheetnames or 'Hoja1' in wb.sheetnames:
            ws_datos = wb['Datos'] if 'Datos' in wb.sheetnames else wb['Hoja1']
            error_fill_cell = PatternFill(start_color="FFCCCC", end_color="FFCCCC", fill_type="solid")
            
            for fila, col_letter, campo, val_orig, val_norm, estado, cod, msg, accion in errores:
                try:
                    col_idx = openpyxl.utils.column_index_from_string(col_letter)
                    cell = ws_datos.cell(row=fila, column=col_idx)
                    cell.fill = error_fill_cell
                    # Agregar comentario con el error
                    from openpyxl.comments import Comment
                    cell.comment = Comment(f"{cod}: {msg}", "Validador")
                except:
                    pass
        
        wb.save(archivo_salida)
        wb.close()
        
        print(f"[OK] Archivo con errores generado: {archivo_salida}")
        print(f"     Errores: {len(errores)}")
        return True
        
    except Exception as e:
        print(f"[ERROR] {e}")
        import traceback
        traceback.print_exc()
        return False

def main():
    if len(sys.argv) < 4:
        print("Uso: python generar_errores_excel.py <resultado_id> <archivo_original.xlsx> <archivo_salida.xlsx>")
        sys.exit(1)
    
    resultado_id = int(sys.argv[1])
    archivo_original = sys.argv[2]
    archivo_salida = sys.argv[3]
    
    if not Path(archivo_original).exists():
        print(f"[ERROR] Archivo original no existe: {archivo_original}")
        sys.exit(1)
    
    exito = generar_archivo_con_errores(resultado_id, archivo_original, archivo_salida)
    sys.exit(0 if exito else 1)

if __name__ == "__main__":
    main()