#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
detectar_columnas.py
Lee un archivo Excel y detecta las columnas con datos en una fila específica.
Retorna lista de columnas con su letra y encabezado.

Uso: python detectar_columnas.py <archivo.xlsx> <hoja> <fila_detectar>
Output: JSON con columnas detectadas
"""

import sys
import json
from pathlib import Path
import openpyxl
from openpyxl.utils import get_column_letter

def detectar_columnas(archivo_path, hoja, fila_detectar):
    """Detecta columnas con datos en la fila especificada"""
    try:
        wb = openpyxl.load_workbook(archivo_path, read_only=True, data_only=True)
        ws = wb[hoja] if hoja in wb.sheetnames else wb.active

        columnas = []
        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=fila_detectar, column=col_idx)
            if cell.value and isinstance(cell.value, str) and cell.value.strip():
                columnas.append({
                    'columna': get_column_letter(col_idx),
                    'encabezado': cell.value.strip()
                })

        wb.close()
        return columnas

    except Exception as e:
        print(f"[ERROR] {e}", file=sys.stderr)
        return []

def main():
    if len(sys.argv) < 4:
        print("Uso: python detectar_columnas.py <archivo.xlsx> <hoja> <fila_detectar>")
        sys.exit(1)

    archivo_path = sys.argv[1]
    hoja = sys.argv[2]
    fila_detectar = int(sys.argv[3])

    if not Path(archivo_path).exists():
        print(f"[ERROR] Archivo no encontrado: {archivo_path}", file=sys.stderr)
        sys.exit(1)

    columnas = detectar_columnas(archivo_path, hoja, fila_detectar)

    # Output JSON para que 4GL capture
    output = {
        'total': len(columnas),
        'columnas': columnas
    }
    print(json.dumps(output, ensure_ascii=False, indent=2))

    sys.exit(0 if columnas else 1)

if __name__ == "__main__":
    main()