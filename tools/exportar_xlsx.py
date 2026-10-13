#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
exportar_xlsx.py
Exporta resultados de validación a Excel (.xlsx) con openpyxl
Uso: python exportar_xlsx.py <resultado_id> <ruta_salida.xlsx>
"""

import sys
import os
import json
from pathlib import Path
from datetime import datetime

# Agregar path para conexión BD
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, numbers
    from openpyxl.utils import get_column_letter
except ImportError:
    print("ERROR: openpyxl no instalado. pip install openpyxl")
    sys.exit(1)

# Mock de conexión BD - en producción usar informixdb o similar
def get_datos_validos(resultado_id):
    """
    En producción: conectar a Informix y ejecutar:
    SELECT * FROM vw_validacion_datos_pivot WHERE resultado_id = ?
    Retorna: lista de dicts {campo: valor} por fila
    """
    # MOCK DATA para testing
    return [
        {'cliente.tipo_doc': 'DNI', 'cliente.num_doc': '30123456', 'cliente.apellidos': 'PEREZ',
         'cliente.nombres': 'JUAN', 'cliente.fecha_nac': '15/03/1985', 'cliente.email': 'juan@email.com'},
        {'cliente.tipo_doc': 'DNI', 'cliente.num_doc': '30789012', 'cliente.apellidos': 'GARCIA',
         'cliente.nombres': 'MARIA', 'cliente.fecha_nac': '22/07/1990', 'cliente.email': 'maria@email.com'},
    ]

def get_config_campos(resultado_id):
    """Obtiene configuración de campos desde BD"""
    # MOCK
    return [
        {'nombre_campo': 'cliente.tipo_doc', 'encabezado': 'Tipo Documento', 'formato_excel': 'General', 'orden': 1},
        {'nombre_campo': 'cliente.num_doc', 'encabezado': 'Número Documento', 'formato_excel': 'General', 'orden': 2},
        {'nombre_campo': 'cliente.apellidos', 'encabezado': 'Apellidos', 'formato_excel': 'General', 'orden': 3},
        {'nombre_campo': 'cliente.nombres', 'encabezado': 'Nombres', 'formato_excel': 'General', 'orden': 4},
        {'nombre_campo': 'cliente.fecha_nac', 'encabezado': 'Fecha Nacimiento', 'formato_excel': 'DD/MM/YYYY', 'orden': 5},
        {'nombre_campo': 'cliente.email', 'encabezado': 'Email', 'formato_excel': '@', 'orden': 6},
    ]

def get_info_validacion(resultado_id):
    """Info cabecera"""
    return {
        'plantilla_nombre': 'Clientes Banco Nación',
        'banco_nombre': 'Banco de la Nación Argentina',
        'fecha_validacion': datetime.now().strftime('%d/%m/%Y %H:%M'),
        'total_filas': 2,
        'filas_ok': 2,
        'filas_error': 0
    }

def exportar_xlsx(resultado_id, ruta_salida):
    """Función principal"""
    try:
        print(f"[INFO] Exportando resultado {resultado_id} a {ruta_salida}")
        
        # Obtener datos
        datos = get_datos_validos(resultado_id)
        config = get_config_campos(resultado_id)
        info = get_info_validacion(resultado_id)
        
        if not datos:
            print("[WARNING] No hay datos válidos para exportar")
            return False
        
        # Crear workbook
        wb = openpyxl.Workbook()
        
        # === HOJA 1: DATOS VÁLIDOS ===
        ws = wb.active
        ws.title = "Datos Validos"
        
        # Estilos
        header_font = Font(bold=True, color="FFFFFF", size=11)
        header_fill = PatternFill(start_color="003366", end_color="003366", fill_type="solid")
        header_align = Alignment(horizontal="center", wrap_text=True)
        thin_border = Border(
            left=Side(style='thin'), right=Side(style='thin'),
            top=Side(style='thin'), bottom=Side(style='thin')
        )
        title_font = Font(bold=True, size=14, color="003366")
        subtitle_font = Font(size=10, color="666666")
        
        # Cabecera con info
        ws.merge_cells('A1:F1')
        ws['A1'] = info['banco_nombre']
        ws['A1'].font = title_font
        ws['A1'].alignment = Alignment(horizontal="center")
        
        ws.merge_cells('A2:F2')
        ws['A2'] = f"Validación: {info['plantilla_nombre']}"
        ws['A2'].font = subtitle_font
        ws['A2'].alignment = Alignment(horizontal="center")
        
        ws.merge_cells('A3:F3')
        ws['A3'] = f"Fecha: {info['fecha_validacion']} | Total: {info['total_filas']} | OK: {info['filas_ok']} | Errores: {info['filas_error']}"
        ws['A3'].font = subtitle_font
        ws['A3'].alignment = Alignment(horizontal="center")
        
        # Fila 5: Encabezados
        row_header = 5
        for idx, campo_cfg in enumerate(config, 1):
            cell = ws.cell(row=row_header, column=idx, value=campo_cfg['encabezado'])
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align
            cell.border = thin_border
        
        # Datos desde fila 6
        for row_idx, fila_datos in enumerate(datos, row_header + 1):
            for col_idx, campo_cfg in enumerate(config, 1):
                nombre_campo = campo_cfg['nombre_campo']
                valor = fila_datos.get(nombre_campo, '')
                cell = ws.cell(row=row_idx, column=col_idx, value=valor)
                cell.border = thin_border
                cell.alignment = Alignment(horizontal="left")
                
                # Formato según config
                fmt = campo_cfg.get('formato_excel', 'General')
                if fmt == 'DD/MM/YYYY':
                    cell.number_format = 'DD/MM/YYYY'
                elif fmt in ('#,##0.00', '0.00'):
                    cell.number_format = '#,##0.00'
                elif fmt == '@':
                    cell.number_format = '@'
        
        # Ajustar anchos
        for idx, campo_cfg in enumerate(config, 1):
            col_letter = get_column_letter(idx)
            ancho = max(len(campo_cfg['encabezado']) + 4, 15)
            ws.column_dimensions[col_letter].width = min(ancho, 40)
        
        # Auto-filtro
        max_col = get_column_letter(len(config))
        ws.auto_filter.ref = f"A{row_header}:{max_col}{row_header + len(datos)}"
        ws.freeze_panes = f"A{row_header + 1}"
        
        # === HOJA 2: RESUMEN ===
        ws2 = wb.create_sheet("Resumen")
        ws2['A1'] = "Resumen de Validación"
        ws2['A1'].font = title_font
        ws2['A3'] = "Banco:"
        ws2['B3'] = info['banco_nombre']
        ws2['A4'] = "Plantilla:"
        ws2['B4'] = info['plantilla_nombre']
        ws2['A5'] = "Fecha Validación:"
        ws2['B5'] = info['fecha_validacion']
        ws2['A6'] = "Total Filas:"
        ws2['B6'] = info['total_filas']
        ws2['A7'] = "Filas Válidas:"
        ws2['B7'] = info['filas_ok']
        ws2['A8'] = "Filas con Error:"
        ws2['B8'] = info['filas_error']
        
        # Estilo resumen
        for row in range(3, 9):
            ws2.cell(row=row, column=1).font = Font(bold=True)
            ws2.column_dimensions['A'].width = 20
            ws2.column_dimensions['B'].width = 40
        
        # === HOJA 3: ERRORES (si los hay) ===
        # En producción: obtener de validacion_detalle donde estado='ERROR'
        # if info['filas_error'] > 0:
        #     ws3 = wb.create_sheet("Errores")
        #     ...
        
        # Guardar
        Path(ruta_salida).parent.mkdir(parents=True, exist_ok=True)
        wb.save(ruta_salida)
        wb.close()
        
        print(f"[OK] Exportado exitosamente: {ruta_salida}")
        return True
        
    except Exception as e:
        print(f"[ERROR] {e}")
        import traceback
        traceback.print_exc()
        return False

def main():
    if len(sys.argv) < 3:
        print("Uso: python exportar_xlsx.py <resultado_id> <ruta_salida.xlsx>")
        sys.exit(1)
    
    resultado_id = int(sys.argv[1])
    ruta_salida = sys.argv[2]
    
    exito = exportar_xlsx(resultado_id, ruta_salida)
    sys.exit(0 if exito else 1)

if __name__ == "__main__":
    main()